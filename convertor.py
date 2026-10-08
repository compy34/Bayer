import pandas as pd
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from act_generator import generate_act_sheet


def build_report_dataframe(df):
    """Build the pharmacy/brand summary used by both the preview and Excel export."""
    if df is None or df.empty:
        return None
    pivot_df = df.pivot_table(
        index='PHARMACY_NAME',
        columns='BRAND',
        values='MERCH',
        aggfunc='sum',
        fill_value=0,
    ).reset_index()
    price_dict = (
        df.drop_duplicates(subset=['BRAND']).set_index('BRAND')['Price'].to_dict()
    )
    brand_columns = pivot_df.columns[1:]
    total_sum = 0
    for brand in brand_columns:
        total_sum += pivot_df[brand] * price_dict.get(brand, 0)
    pivot_df['Загальна сума'] = total_sum
    return pivot_df


def generate_excel_report(
    df,
    df_req,
    output_filename='marketing_report.xlsx',
    period_start=None,
    period_end=None,
    selected_pharmacy='Усі аптеки',
):
    if df is None or df.empty:
        return

    pivot_df = build_report_dataframe(df)
    brand_columns = pivot_df.columns[1:]

    with pd.ExcelWriter(output_filename, engine='openpyxl') as writer:
        pivot_df.to_excel(writer, index=False, startrow=9, sheet_name='Звіт')

    import openpyxl
    wb = openpyxl.load_workbook(output_filename)
    ws = wb['Звіт']

    ws.delete_rows(10, 2)

    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000'),
    )
    center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    right_alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)

    num_brands = len(brand_columns)
    sum_col_idx = 3 + num_brands
    header_start_col = 3
    header_end_col = max(3 + num_brands - 1, 8)

    start_let = get_column_letter(header_start_col)
    end_let = get_column_letter(header_end_col)

    contract_no = '260326'
    contract_date = '01 березня 2026 року'
    executor_name = 'ПрАТ «Аптеки Запоріжжя»'

    # Універсальний пошук реквізитів за chain_name або Клиент из базы данных
    if df_req is not None and not df_req.empty:
        matched_row = None
        clean_selected = str(selected_pharmacy).strip()

        target_col = None
        for col_candidate in [
            'chain_name',
            'Клиент из базы данных',
            'Мережа',
            'Назва',
        ]:
            if col_candidate in df_req.columns:
                target_col = col_candidate
                break

        if clean_selected and clean_selected != 'Усі аптеки' and target_col:
            filtered_req = df_req[
                df_req[target_col].astype(str).str.strip().str.lower()
                == clean_selected.lower()
            ]
            if not filtered_req.empty:
                matched_row = filtered_req.iloc[0]

        if matched_row is None:
            matched_row = df_req.iloc[0]

        executor_name = str(
            matched_row.get(
                'Виконавець', matched_row.get('Юридическое лицо', executor_name)
            )
        )
        contract_no = str(matched_row.get('№ договора', contract_no))
        contract_date = str(matched_row.get('Дата договора', contract_date))

    period_start = period_start or '01.08.2026'
    period_end = period_end or '31.08.2026'
    period_str = f'за період з "{period_start}" по "{period_end}".'
    end_date_str = period_end

    ws.merge_cells(f'{start_let}1:{end_let}1')
    ws[f'{start_let}1'] = 'ДОДАТКОВА УГОДА № ________'
    ws[f'{start_let}1'].font = Font(bold=True, size=10)
    ws[f'{start_let}1'].alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells(f'{start_let}2:{end_let}2')
    ws[f'{start_let}2'] = (
        f'до договору про надання послуг № {contract_no} від '
        f'{contract_date} (Маркетингові послуги)'
    )
    ws[f'{start_let}2'].font = Font(bold=True, size=10)
    ws[f'{start_let}2'].alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells(f'{start_let}3:{end_let}3')
    ws[f'{start_let}3'] = 'м. Київ'
    ws[f'{start_let}3'].font = Font(bold=True, size=10)
    ws[f'{start_let}3'].alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells(f'{start_let}4:{end_let}4')
    ws[f'{start_let}4'] = f'{period_str}'
    ws[f'{start_let}4'].font = Font(bold=True, size=10)
    ws[f'{start_let}4'].alignment = Alignment(horizontal='center', vertical='center')

    ws.cell(row=6, column=header_start_col, value='Дата').font = Font(size=10)
    ws.cell(row=6, column=header_start_col).alignment = left_alignment
    ws.cell(row=6, column=header_end_col, value=end_date_str).font = Font(size=10)
    ws.cell(row=6, column=header_end_col).alignment = right_alignment

    ws.merge_cells(start_row=8, start_column=1, end_row=8, end_column=sum_col_idx)
    ws.cell(
        row=8,
        column=1,
        value=(
            f'{executor_name} в особі представника діє на підставі статуту, '
            'та Замовник уклали цю додаткову угоду про наступне:\n1. '
            'Сторони домовилися затвердити Маркетинговий звіт за звітний період.'
        ),
    )
    ws.cell(row=8, column=1).font = Font(size=10)
    ws.cell(row=8, column=1).alignment = Alignment(
        horizontal='left', vertical='center', wrap_text=True
    )
    ws.row_dimensions[8].height = 45

    header_row_1 = 10
    header_row_2 = 11

    ws.merge_cells(
        start_row=header_row_1, start_column=1, end_row=header_row_2, end_column=1
    )
    ws.cell(row=header_row_1, column=1, value='№ з/п')

    ws.merge_cells(
        start_row=header_row_1, start_column=2, end_row=header_row_2, end_column=2
    )
    ws.cell(
        row=header_row_1,
        column=2,
        value=('Найменування аптечного закладу (Аптеки / Мережі)'),
    )

    start_col = 3
    end_col = start_col + num_brands - 1
    start_col_letter = get_column_letter(start_col)
    end_col_letter = get_column_letter(end_col)

    ws.merge_cells(f'{start_col_letter}{header_row_1}:{end_col_letter}{header_row_1}')
    ws[f'{start_col_letter}{header_row_1}'] = 'Бренди'

    for i, brand in enumerate(brand_columns):
        col_idx = start_col + i
        ws.cell(row=header_row_2, column=col_idx, value=brand)

    sum_col_letter = get_column_letter(sum_col_idx)
    ws.merge_cells(
        start_row=header_row_1,
        start_column=sum_col_idx,
        end_row=header_row_2,
        end_column=sum_col_idx,
    )
    ws.cell(row=header_row_1, column=sum_col_idx, value='Загальна сума')

    ws.row_dimensions[header_row_1].height = 30
    ws.row_dimensions[header_row_2].height = 25

    start_data_row = 12
    last_data_row = start_data_row + len(pivot_df) - 1

    for idx, row_data in enumerate(pivot_df.iterrows(), start=start_data_row):
        row_values = row_data[1]
        cell_no = ws.cell(row=idx, column=1, value=idx - start_data_row + 1)
        cell_no.alignment = center_alignment
        cell_no.border = thin_border

        cell_name = ws.cell(row=idx, column=2, value=row_values['PHARMACY_NAME'])
        cell_name.alignment = left_alignment
        cell_name.border = thin_border

        for i, brand in enumerate(brand_columns):
            cell_brand = ws.cell(
                row=idx, column=start_col + i, value=row_values[brand]
            )
            cell_brand.alignment = center_alignment
            cell_brand.border = thin_border

        cell_sum = ws.cell(
            row=idx, column=sum_col_idx, value=row_values['Загальна сума']
        )
        cell_sum.alignment = center_alignment
        cell_sum.border = thin_border
        cell_sum.number_format = '#,##0.00'

    for row in ws.iter_rows(
        min_row=header_row_1,
        max_row=header_row_2,
        min_col=1,
        max_col=sum_col_idx,
    ):
        for cell in row:
            cell.alignment = center_alignment
            cell.border = thin_border
            cell.font = Font(bold=True, size=9)

    start_row_bottom = last_data_row + 2
    col_label = sum_col_idx - 1
    col_value = sum_col_idx

    ws.cell(
        row=start_row_bottom, column=col_label, value='Всього без ПДВ'
    ).alignment = right_alignment
    ws.cell(row=start_row_bottom, column=col_label).font = Font(bold=True)
    ws.cell(row=start_row_bottom, column=col_label).border = thin_border

    val1 = ws.cell(
        row=start_row_bottom,
        column=col_value,
        value=f'=SUM({sum_col_letter}{start_data_row}:{sum_col_letter}{last_data_row})',
    )
    val1.alignment = right_alignment
    val1.font = Font(bold=True)
    val1.border = thin_border
    val1.number_format = '#,##0.00'

    row_vat = start_row_bottom + 1
    ws.cell(row=row_vat, column=col_label, value='ПДВ 20%').alignment = right_alignment
    ws.cell(row=row_vat, column=col_label).font = Font(bold=True)
    ws.cell(row=row_vat, column=col_label).border = thin_border

    val2 = ws.cell(
        row=row_vat,
        column=col_value,
        value=f'={sum_col_letter}{start_row_bottom} * 0.2',
    )
    val2.alignment = right_alignment
    val2.font = Font(bold=True)
    val2.border = thin_border
    val2.number_format = '#,##0.00'

    row_total = start_row_bottom + 2
    ws.cell(
        row=row_total, column=col_label, value='Всього з ПДВ'
    ).alignment = right_alignment
    ws.cell(row=row_total, column=col_label).font = Font(bold=True)
    ws.cell(row=row_total, column=col_label).border = thin_border

    val3 = ws.cell(
        row=row_total,
        column=col_value,
        value=(
            f'={sum_col_letter}{start_row_bottom} +'
            f' {sum_col_letter}{row_vat}'
        ),
    )
    val3.alignment = right_alignment
    val3.font = Font(bold=True)
    val3.border = thin_border
    val3.number_format = '#,##0.00'

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        if start_col <= col[0].column < sum_col_idx:
            ws.column_dimensions[col_letter].width = 16
        elif col[0].column == 1:
            ws.column_dimensions[col_letter].width = 8
        elif col[0].column == 2:
            ws.column_dimensions[col_letter].width = 50
        else:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 15)

    total_sum_no_vat = pivot_df['Загальна сума'].sum()
    generate_act_sheet(
        wb=wb,
        df_req=df_req,
        total_sum_without_vat=total_sum_no_vat,
        period_start=period_start,
        period_end=period_end,
        selected_pharmacy=selected_pharmacy,
        act_number='1',
    )

    wb.save(output_filename)
    return output_filename